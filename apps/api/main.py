import base64
from security.authorization import AuthorizationError
from security.pii_filter import anonymize_pii
from security.prompt_guard import (
    PromptInjectionError,
    validate_user_prompt,
)
from security.output_filter import redact_secrets
from contextlib import asynccontextmanager
from io import BytesIO

from fastapi import (
    FastAPI,
    HTTPException,
)

from openpyxl import load_workbook
from pydantic import BaseModel

from orchestration.delivery_workflow import (
    run_delivery_workflow,
)

from retrieval.hybrid_rag import (
    initialize_hybrid_rag,
)


# ---------------------------------------------------------
# FastAPI lifespan
# ---------------------------------------------------------

@asynccontextmanager
async def lifespan(
    app: FastAPI,
):

    print(
        "[Startup] Warming Hybrid RAG..."
    )

    initialize_hybrid_rag()

    print(
        "[Startup] Hybrid RAG ready."
    )

    yield

    print(
        "[Shutdown] "
        "MAQ Delivery Agent stopped."
    )


# ---------------------------------------------------------
# FastAPI application
# ---------------------------------------------------------

app = FastAPI(
    title=(
        "MAQ Intelligent "
        "Client Delivery Agent"
    ),
    version="0.5.0",
    lifespan=lifespan,
)


# ---------------------------------------------------------
# Request model
# ---------------------------------------------------------

class DeliveryQueryRequest(
    BaseModel
):

    user_question: str
    user_id: str
    file_name: str
    file_content_base64: str


# ---------------------------------------------------------
# Health endpoint
# ---------------------------------------------------------

@app.get("/health")
async def health():

    return {
        "status": "ok",
        "service": (
            "maq-client-delivery-agent"
        ),
        "environment": "dev",
        "hybrid_rag": "ready",
        "architecture": "multi-agent",
        "agents": 3,
    }


# ---------------------------------------------------------
# SharePoint workbook parser
# ---------------------------------------------------------

def parse_project_register(
    file_content_base64: str,
) -> list[dict]:
    """
    Parse the SharePoint-hosted Excel workbook
    and return MAQProjectRegister records.
    """

    workbook_bytes = (
        base64.b64decode(
            file_content_base64
        )
    )

    workbook = load_workbook(
        BytesIO(
            workbook_bytes
        ),
        data_only=True,
    )

    target_sheet = None
    target_table = None

    for sheet in workbook.worksheets:

        for table in (
            sheet.tables.values()
        ):

            if (
                table.name
                == "MAQProjectRegister"
            ):

                target_sheet = sheet
                target_table = table

                break

        if target_table:
            break

    if (
        target_sheet is None
        or target_table is None
    ):

        raise ValueError(
            "MAQProjectRegister table "
            "was not found."
        )

    cells = target_sheet[
        target_table.ref
    ]

    headers = [
        cell.value
        for cell in cells[0]
    ]

    projects = []

    for row in cells[1:]:

        values = [
            cell.value
            for cell in row
        ]

        project = dict(
            zip(
                headers,
                values,
            )
        )

        for key, value in (
            project.items()
        ):

            if hasattr(
                value,
                "isoformat",
            ):

                project[key] = (
                    value.isoformat()
                )

        projects.append(
            project
        )

    return projects


# ---------------------------------------------------------
# SharePoint debug endpoint
# ---------------------------------------------------------

@app.post(
    "/sharepoint/project-register"
)
async def read_project_register(
    request: DeliveryQueryRequest,
):

    try:

        projects = (
            parse_project_register(
                request.file_content_base64
            )
        )

        return {
            "success": True,
            "source": "SharePoint",
            "user_question":
                request.user_question,
            "user_id":
                request.user_id,
            "file_name":
                request.file_name,
            "table_name":
                "MAQProjectRegister",
            "project_count":
                len(projects),
            "projects":
                projects,
        }

    except Exception as exc:

        print(
            "[SharePointParser] "
            f"{type(exc).__name__}: "
            f"{exc}"
        )

        raise HTTPException(
            status_code=500,
            detail=(
                "The SharePoint project "
                "register could not be parsed."
            ),
        )


# ---------------------------------------------------------
# Main multi-agent delivery endpoint
# ---------------------------------------------------------

@app.post("/delivery/query")
async def delivery_query(
    request: DeliveryQueryRequest,
):
    """
    Executes the MAQ three-agent workflow:

    Agent 1:
        MAQPortfolioEvidenceAgent

    Agent 2:
        MAQEngineeringEvidenceAgent

    Agent 3:
        MAQDeliveryAnalystAgent

    Portfolio and Engineering agents execute
    concurrently. Their evidence is validated
    and passed to the Analyst Agent.
    """

    try:

        # -------------------------------------------------
        # Parse live SharePoint project register
        # -------------------------------------------------

        projects = (
            parse_project_register(
                request.file_content_base64
            )
        )

        print(
            "[DeliveryQuery] "
            f"Loaded {len(projects)} "
            "SharePoint projects."
        )


        # -------------------------------------------------
        # Request-level source tracking
        # -------------------------------------------------

        sources_used = set()

        source_order = [
            "SharePoint",
            "Dataverse",
            "Azure DevOps",
            "MAQ Delivery Knowledge",
        ]


        def mark_source(
            source_name: str,
        ) -> None:

            sources_used.add(
                source_name
            )

            print(
                "[SourceTracking] "
                f"Used: {source_name}"
            )


        # -------------------------------------------------
        # PII detection and masking
        # -------------------------------------------------

        pii_result = anonymize_pii(
            request.user_question
        )

        sanitized_question = (
            pii_result["sanitized_text"]
        )

        if pii_result["pii_detected"]:
            print(
                "[Security] PII detected and masked:",
                [
                    entity["entity_type"]
                    for entity
                    in pii_result["entities"]
                ],
            )


        # -------------------------------------------------
        # Prompt-injection guard
        # -------------------------------------------------

        validate_user_prompt(
            sanitized_question
        )

        print(
            "[Security] Prompt guard passed."
        )


        # -------------------------------------------------
        # Run 3-agent orchestration
        # -------------------------------------------------

        print(
            "[DeliveryQuery] Starting "
            "multi-agent workflow..."
        )

        workflow_result = (
            await run_delivery_workflow(
                user_id=request.user_id,
                user_question=sanitized_question,
                projects=projects,
                mark_source=mark_source,
            )
        )

        safe_answer, secret_detected = redact_secrets(
            workflow_result["answer"]
        )

        if secret_detected:
            print(
                "[Security] Secret-like content redacted from output."
            )

        workflow_result["answer"] = safe_answer


        # -------------------------------------------------
        # Actual evidence sources used
        # -------------------------------------------------

        actual_sources = [
            source
            for source in source_order
            if source in sources_used
        ]

        print(
            "[SourceTracking] "
            "Final sources:",
            actual_sources,
        )


        # -------------------------------------------------
        # Final HTTP response
        # -------------------------------------------------

        return {
            "success": True,

            "user_id":
                request.user_id,

            "question":
                sanitized_question,

            "answer":
                workflow_result["answer"],

            "sources":
                actual_sources,

            "workflow":
                workflow_result[
                    "workflow"
                ],
        }


    except PromptInjectionError as exc:

        print(
            "[Security] Prompt injection blocked:",
            str(exc),
        )

        raise HTTPException(
            status_code=400,
            detail=(
                "The request was blocked by "
                "the prompt security policy."
            ),
        )

    except AuthorizationError as exc:

        print(
            "[DeliveryQuery] Authorization denied:",
            str(exc),
        )

        raise HTTPException(
            status_code=403,
            detail="You are not authorized to access the requested delivery information.",
        )

    except Exception as exc:

        print(
            "[DeliveryQuery] Error: "
            f"{type(exc).__name__}: "
            f"{exc}"
        )

        raise HTTPException(
            status_code=500,
            detail=(
                "Delivery data could not "
                "be retrieved."
            ),
        )